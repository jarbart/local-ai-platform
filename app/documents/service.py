from hashlib import sha256
from pathlib import Path

import pymupdf
from docx import Document

from app.documents.models import NormalizedDocument


class DocumentService:
    def extract_text(
        self,
        file_path: Path,
        filename: str,
        content_type: str,
    ) -> NormalizedDocument:
        document_id = self._calculate_document_id(file_path)

        if content_type == "text/plain":
            return self._extract_text_file(
                file_path=file_path,
                document_id=document_id,
                filename=filename,
                content_type=content_type,
            )

        if content_type == "application/pdf":
            return self._extract_pdf(
                file_path=file_path,
                document_id=document_id,
                filename=filename,
                content_type=content_type,
            )

        if content_type == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            return self._extract_docx(
                file_path=file_path,
                document_id=document_id,
                filename=filename,
                content_type=content_type,
            )

        raise ValueError(
            f"Unsupported content type: {content_type}"
        )

    def _extract_text_file(
        self,
        file_path: Path,
        document_id: str,
        filename: str,
        content_type: str,
    ) -> NormalizedDocument:
        text = file_path.read_text(
            encoding="utf-8",
        )

        return NormalizedDocument(
            document_id=document_id,
            filename=filename,
            content_type=content_type,
            text=text,
            metadata={
                "page_count": 1,
                "pages": [
                    {
                        "page_number": 1,
                        "text": text,
                    }
                ],
                "character_count": len(text),
                "extraction_method": "plain_text",
            },
        )

    def _extract_pdf(
        self,
        file_path: Path,
        document_id: str,
        filename: str,
        content_type: str,
    ) -> NormalizedDocument:
        with pymupdf.open(file_path) as document:
            pages = [
                {
                    "page_number": page_number,
                    "text": page.get_text(),
                }
                for page_number, page in enumerate(
                    document,
                    start=1,
                )
            ]

        text = "\n".join(
            page["text"]
            for page in pages
        )

        return NormalizedDocument(
            document_id=document_id,
            filename=filename,
            content_type=content_type,
            text=text,
            metadata={
                "page_count": len(pages),
                "pages": pages,
                "character_count": len(text),
                "extraction_method": "pymupdf",
            },
        )

    def _extract_docx(
        self,
        file_path: Path,
        document_id: str,
        filename: str,
        content_type: str,
    ) -> NormalizedDocument:
        document = Document(file_path)

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        text = "\n".join(paragraphs)

        return NormalizedDocument(
            document_id=document_id,
            filename=filename,
            content_type=content_type,
            text=text,
            metadata={
                "page_count": 1,
                "pages": [
                    {
                        "page_number": 1,
                        "text": text,
                    }
                ],
                "character_count": len(text),
                "extraction_method": "python_docx",
            },
        )

    @staticmethod
    def _calculate_document_id(file_path: Path) -> str:
        hasher = sha256()

        with file_path.open("rb") as file:
            for chunk in iter(
                lambda: file.read(1024 * 1024),
                b"",
            ):
                hasher.update(chunk)

        return hasher.hexdigest()