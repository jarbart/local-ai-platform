from pathlib import Path

from app.documents.models import NormalizedDocument
from app.documents.service import DocumentService

from docx import Document


def test_extract_text_from_pdf():
    service = DocumentService()

    document = service.extract_text(
        Path("test.pdf"),
        filename="test.pdf",
        content_type="application/pdf",
    )

    assert isinstance(document, NormalizedDocument)
    assert document.filename == "test.pdf"
    assert document.content_type == "application/pdf"
    assert document.document_id
    assert document.text.strip()

    assert document.metadata["page_count"] > 0
    assert len(document.metadata["pages"]) == document.metadata["page_count"]

    first_page = document.metadata["pages"][0]

    assert first_page["page_number"] == 1
    assert first_page["text"]


def test_extract_text_file():
    file_path = Path("test.txt")

    service = DocumentService()

    document = service.extract_text(
        file_path=file_path,
        filename="test.txt",
        content_type="text/plain",
    )

    assert isinstance(document, NormalizedDocument)
    assert document.filename == "test.txt"
    assert document.content_type == "text/plain"
    assert document.document_id
    assert document.text.strip()

    assert document.metadata["page_count"] == 1
    assert document.metadata["extraction_method"] == "plain_text"
    assert document.metadata["pages"][0]["page_number"] == 1


def test_extract_text_from_docx(tmp_path):
    file_path = tmp_path / "test.docx"

    document = Document()
    document.add_paragraph("Mój dokument DOCX.")
    document.add_paragraph(
        "Zawiera informacje o lokalnej platformie AI."
    )
    document.save(file_path)

    service = DocumentService()

    result = service.extract_text(
        file_path=file_path,
        filename="test.docx",
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )

    assert isinstance(result, NormalizedDocument)
    assert result.filename == "test.docx"
    assert result.content_type == (
        "application/"
        "vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    assert result.document_id
    assert "Mój dokument DOCX." in result.text
    assert "lokalnej platformie AI" in result.text

    assert result.metadata["page_count"] == 1
    assert result.metadata["extraction_method"] == "python_docx"
    assert result.metadata["pages"][0]["page_number"] == 1